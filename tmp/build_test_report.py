from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / 'KET-QUA-TEST_A313-V.06'
OUT = ROOT / 'BAO-CAO_KIEM-TRA_A313-V.06.docx'
d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
sec.left_margin, sec.right_margin = Cm(2.1), Cm(1.8)
sec.footer_distance = Cm(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st = d.styles[name]
    st.font.name = 'Times New Roman'
    st.font.color.rgb = RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
normal = d.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(6)
for name, size in [('Title',23),('Subtitle',13),('Heading 1',16),('Heading 2',12)]:
    d.styles[name].font.size=Pt(size)
    d.styles[name].paragraph_format.space_after=Pt(8)
d.styles['Heading 1'].paragraph_format.space_before=Pt(0)
d.styles['Heading 2'].paragraph_format.space_before=Pt(8)
d.styles['Caption'].font.size=Pt(9)
d.styles['Caption'].font.italic=True
d.styles['Caption'].paragraph_format.space_after=Pt(7)
foot = sec.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('A313-V.06  |  Trang '); r.font.size=Pt(9)
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); foot._p.append(f)
d.core_properties.title='Báo cáo nghiên cứu và kiểm tra chức năng tấm A313-V.06'
d.core_properties.subject='Khảo sát sơ đồ và kết quả thử nghiệm chức năng trên bàn'
d.core_properties.author=''
d.core_properties.keywords='A313-V.06, 400 Hz, kiểm tra chức năng, 16 kênh'

def p(text, bold=False):
    x=d.add_paragraph(); x.add_run(text).bold=bold; return x
def h(text): d.add_heading(text,level=1)
def sub(text): d.add_heading(text,level=2)
def page(): d.add_page_break()
def table(headers, rows, widths):
    t=d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for i,txt in enumerate(headers): t.rows[0].cells[i].text=txt
    for vals in rows:
        cells=t.add_row().cells
        for i,txt in enumerate(vals): cells[i].text=str(txt)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        ns=OxmlElement('w:cantSplit'); trpr.append(ns)
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE6F1' if ri==0 else ('F4F6F8' if ri%2==0 else 'FFFFFF')); pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'75'); el.set(qn('w:type'),'dxa'); margins.append(el)
            pr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            pr.append(borders)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2); para.paragraph_format.space_before=Pt(2)
                para.paragraph_format.line_spacing=1.0
                for run in para.runs:
                    run.font.size=Pt(10); run.bold=(ri==0)
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def photo(name,width,caption,base=IMG):
    x=d.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER
    x.paragraph_format.space_after=Pt(3); x.paragraph_format.keep_with_next=True
    shape=x.add_run().add_picture(str(base/name),width=Cm(width))
    shape._inline.docPr.set('descr',caption)
    c=d.add_paragraph(caption,style='Caption'); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
def pair(a,b,caption):
    x=d.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER
    x.paragraph_format.space_after=Pt(3); x.paragraph_format.keep_with_next=True
    for i,name in enumerate([a,b]):
        if i: x.add_run('  ')
        shape=x.add_run().add_picture(str(IMG/name),width=Cm(8.1))
        shape._inline.docPr.set('descr',caption+(' — trái' if i==0 else ' — phải'))
    c=d.add_paragraph(caption,style='Caption'); c.alignment=WD_ALIGN_PARAGRAPH.CENTER

# 1
d.add_paragraph('BÁO CÁO NGHIÊN CỨU\nVÀ KIỂM TRA CHỨC NĂNG',style='Title')
d.add_paragraph('Tấm khuếch đại tín hiệu ra A313-V.06',style='Subtitle')
p('Báo cáo kỹ thuật trình lãnh đạo • Tổng hợp ngày 17 tháng 9 năm 2026')
sub('Kết luận chính')
p('Tấm A313-V.06 được đánh giá hoạt động tốt về chức năng trong phạm vi các phép thử đã thực hiện. Kết quả kiểm tra cho thấy mạch tạo điều khiển 400 Hz làm việc, hai tín hiệu M và M̄ bù trạng thái, đường truyền kênh 1a đáp ứng đúng ở hai pha tĩnh và các kênh cho đầu ra động tương đồng. Khi thay đổi mức và dấu điện áp đầu vào, đầu ra thay đổi phù hợp với nguyên lý điều chế của tấm.',True)
p('Cơ sở đánh giá gồm nghiên cứu chuỗi tín hiệu trong khối A313-V, đối chiếu sơ đồ nguyên lý với kết nối linh kiện, xây dựng quy trình thử theo từng cụm chức năng và kiểm tra thực tế bằng nguồn lập trình, đồng hồ số và máy hiện sóng. Ảnh đo được lưu theo mã bài và điểm đo để phục vụ đối chiếu về sau.')
table(['Nội dung','Kết quả nổi bật'],[
('Bộ điều khiển chung','M và M̄ cùng 400,0 Hz; biên độ đo đồng thời 26,27 V và 26,26 V; bù trạng thái ở phần ổn định.'),
('Đường truyền tín hiệu','B07 xác nhận cùng dấu và đảo dấu ở kênh 1a với Vin = +1 V.'),
('Thử động và thử tải','Đã kiểm tra lần lượt các kênh ở B09 và B10; ảnh đại diện cho mức +1 V và +4 V có Amplitude 1,975 V và 7,957 V.'),
('Kiểm tra mở rộng','B11 và B12 kiểm tra một số điểm, chưa ghi nhận bất thường tại các điểm đã thử.')
],[4,13.1])
sub('Phạm vi kết luận')
p('Kết luận “tấm tốt” trong báo cáo này là kết luận chức năng trên bàn theo điều kiện đã thử. B11–B12 mới kiểm tra chọn mẫu; kết quả không thay thế nghiệm thu đầy đủ toàn bộ tổ hợp tải, thử môi trường hoặc thử tích hợp toàn hệ thống. Chưa phát hiện dấu hiệu buộc phải chỉ định thay thế linh kiện từ các kết quả hiện có.')
p('Người thực hiện: ..............................................................\nĐơn vị: ..............................................................................')

# 2
page(); h('1 Vai trò của tấm trong khối A313 V')
p('Theo Hướng dẫn sử dụng, dữ liệu ra từ bộ nhớ được đưa qua cụm biến đổi A313-V.04 để tạo tín hiệu tương tự. Cụm A313-V.05 lựa chọn và lưu giữ mức điện áp của từng kênh bằng các bộ nhớ tương tự. Nhóm tín hiệu 1a–8a và 1b–8b được đưa tiếp đến A313-V.06 để điều chế thành tín hiệu xoay chiều 400 Hz phục vụ nhánh biến thế sin–cos của hệ thống ngoại vi. [1]')
p('Chuỗi chức năng: dữ liệu số → A313-V.04 biến đổi → A313-V.05 lưu giữ mức DC → A313-V.06 điều chế và đệm dòng → đầu cắm 3Ш3 → hệ thống ngoại vi.',True)
p('Như vậy, tấm xử lý 16 đường tín hiệu tương tự, tổ chức thành 8 cặp a/b. Độ lớn đầu vào quyết định mức đầu ra; dấu đầu vào quyết định sự đảo pha 180° so với cùng trạng thái điều khiển. Việc thử cả điện áp dương và âm vì thế có ý nghĩa kiểm tra thông tin dấu, không chỉ kiểm tra có điện áp ra.')
photo('B00_Mat-linh-kien_DC11_DC12.png',6.7,'Hình 1. Tấm thực tế và bộ gá; hai giắc DC11, DC12 là giao diện cấp nguồn và đo kiểm.')
p('Các liên hệ cần phân biệt: A313-V.04 thực hiện biến đổi; A313-V.05 giữ mức cho từng kênh; A313-V.06 tạo tín hiệu ra 400 Hz. Nguồn trong khối thuộc cụm A313-V.08.01. A313-V.02 tham gia các đường điều khiển chuyển mạch trong hệ thống; ở phép thử độc lập, bộ gá cung cấp trực tiếp nguồn và tín hiệu cần thiết cho tấm .06. [1, 2]')

# 3
page(); h('2 Phân tích sơ đồ và cơ sở xây dựng phép thử')
sub('Cụm tạo tín hiệu điều khiển chung')
p('Tín hiệu tham chiếu vào DC11.A5/B5/C5 là một net chung. U14 và U15 nhận tín hiệu qua hai bộ chia áp khác nhau; U16 chọn các mức ngưỡng, U17–U18 tạo xung và chuyển mức điều khiển. Đầu DC11.A6 có xung khoảng 0/+3,5 V; A7 và A8 đưa ra M và M̄ với hai mức khoảng −23/+3,5 V. Các đầu ra này phải được kiểm tra cả biên độ, tần số và quan hệ trạng thái. [2, 3]')
sub('Một kênh điều chế điển hình')
photo('channel1.png',17,'Hình 2. Trích BV.22: kênh 1a từ khóa tương tự đến OA, transistor đệm và hồi tiếp.',ROOT/'tmp/pdfs/review')
p('Kênh 1a nhận Vin tại DC12.A7. Khóa nối tiếp U2B đưa Vin tới nút S; khóa shunt U1A kéo S về GND ở pha còn lại. R37 dẫn S tới U19.10, còn R85 dẫn Vin tới U19.9. Cặp Q1/Q2 đệm dòng; đầu ra qua R39 đến DC12.A17. R86 và C1 lấy hồi tiếp từ sau R39, vì vậy phép thử có tải kiểm tra được cả tầng đệm và khả năng duy trì điện áp của vòng hồi tiếp.')
p('Với R85 = R86 danh định 10 kΩ, quan hệ xác lập gần đúng là Vo = 2S − Vin. Khi S gần Vin thì Vo gần +Vin; khi S gần 0 thì Vo gần −Vin. Hai phép giữ pha bằng tham chiếu +5 V DC và −5 V DC giúp quan sát riêng từng trạng thái, rồi phép thử 400 Hz kiểm tra sự chuyển đổi liên tục giữa chúng.')
sub('Các chi tiết đã được đối chiếu để tránh chẩn đoán nhầm')
table(['Chi tiết sơ đồ','Ý nghĩa đối với phép kiểm tra'],[
('R24 = 2,05 kΩ; cầu 0 Ω là R124','Giữ đúng hệ số chia áp U15; không nhầm R124 với R63 ở kênh 5a.'),
('U14.9/U15.9 khác U14.5/U15.5','Chân 9 là sin đã chia áp; chân 5 là đầu ra chuyển mức. Phân biệt điểm đo trước khi kết luận lỗi.'),
('Ngưỡng H14 và H15 tách riêng','Không coi hai chân 10 là cùng một nút; quan hệ pha phải kiểm tra theo đúng sơ đồ.'),
('16 kênh và các vỏ khóa dùng chung','Đã lập ánh xạ đủ các kênh; nhóm U9 phục vụ 5a/6a, U8 phục vụ 5b/6b, tránh suy theo thứ tự vị trí.')
],[5.1,12])

# 4
page(); h('3 Phương pháp và điều kiện kiểm tra')
p('Quy trình đi từ kiểm tra nguội và nguồn tới tín hiệu chung, đường truyền tĩnh, đáp ứng động và tải. Cách chia bài giúp phân biệt lỗi chung của tấm với lỗi một kênh, đồng thời kiểm tra từng giả thiết phân tích bằng điểm đo cụ thể thay vì chỉ quan sát đầu ra cuối cùng.')
table(['Nhóm bài','Nội dung kiểm tra','Phạm vi đã ghi nhận'],[
('B00–B03','Nhận dạng, bộ gá, đo nguội, nguồn nuôi','Thông qua theo xác nhận; có ảnh dòng nguồn ban đầu.'),
('B04–B06','Chia áp tham chiếu, xung, M/M̄, giữ pha DC','Có ảnh B04/B05; B06 và phần còn lại B05 được xác nhận đạt.'),
('B07','Đường khóa và hồi tiếp khi giữ hai pha','Kênh 1a, Vin=+1 V; hai pha được xác nhận đúng.'),
('B08','Đặc tuyến DC nhiều mức từng kênh','Chưa có bảng Vin/Vo chi tiết để tính gain, offset từng kênh.'),
('B09–B10','Thử động không tải và có tải','Người thực hiện xác nhận đã đo lần lượt toàn bộ kênh; hai ảnh tải đại diện.'),
('B11–B12','Ảnh hưởng chéo và điều kiện mở rộng','Đã thử một số điểm; chưa ghi nhận vấn đề tại các điểm đó.')
],[2.3,7.0,7.8])
p('Thiết bị sử dụng: Keithley DMM6500; Tektronix MSO44; nguồn bốn kênh Rohde & Schwarz HMP4040; nguồn AC California Instruments 2003RP. Tham chiếu động đặt 5 Vrms/400 Hz. B10 được xây dựng với tải thuần trở 1 kΩ, tối thiểu 0,5 W, thử từng kênh; ảnh không ghi trị số tải thực đo hoặc mã kênh riêng.')
photo('B03_Dong-nguon_HMP4040_CH1-CH4.png',10.6,'Hình 3. Dòng nguồn ban đầu: CH1 112,9 mA; CH2 113,8 mA; CH3 5,4 mA; CH4 5,5 mA.')
p('Bốn rail tại tấm là +12,6 V, −12,6 V, −23 V và +3,5 V. Dòng hai kênh nguồn 12,6 V chênh 0,9 mA; người thực hiện đã xác nhận điện áp tại các điểm cấp đúng. Đây là mốc ban đầu của phép thử, không phải số dòng có tải cực đại.')

# 5
page(); h('4 Kết quả tín hiệu vào và bộ tạo điều khiển')
p('Tại chân 9 U14 và U15, dạng sóng thu được là sin; tại chân 5, dạng sóng chuyển thành xung 400 Hz. Kết quả này phù hợp với cấu trúc chia áp và so sánh ngưỡng, xác nhận được đường đi tín hiệu qua hai khối. Các điểm đo đã được đối chiếu lại để thống nhất tên chân trong hồ sơ ảnh.')
pair('B04_U14_chan-9_IN_Amplitude-1p361V.png','B04_U15_chan-9_IN_Amplitude-1p675V.png','Hình 4. Đầu vào U14.9 bên trái và U15.9 bên phải; đều có dạng sin.')
pair('B04_U14_chan-5_OUT_Amplitude-19p67V_400Hz.png','B04_U15_chan-5_OUT_Amplitude-19p73V_400Hz.png','Hình 5. Đầu ra U14.5 và U15.5; Amplitude 19,67 V và 19,73 V, tần số 400 Hz.')
table(['Điểm','Số đọc trên ảnh','Nhận xét'],[
('U14.9','Amplitude 1,361 V; cursor ΔV ≈ 1,408 V','Sin đã giảm biên độ.'),
('U15.9','Amplitude 1,675 V; cursor ΔV ≈ 1,734 V','Biên độ lớn hơn nhánh U14, đúng chiều quan hệ chia áp.'),
('U14.5 / U15.5','19,67 / 19,73 V; cùng 400 Hz','Có chuyển mức ở đầu ra hai OA.')
],[2.5,8,6.6])
p('Amplitude được giữ đúng tên phép đo của máy, không đổi thành Vrms. Các ảnh chân 9 có cảnh báo Low amplitude ở phép đo tần số nên không dùng số tần số đó để kết luận. Giá trị chia áp danh định khoảng 1,462/1,811 Vpp với sin 5 Vrms là mốc phân tích; các ảnh hiện có xác nhận dạng tín hiệu, chưa dùng để chứng nhận dung sai chia áp.')

# 6
page(); h('5 Kết quả xung và hai tín hiệu bù trạng thái')
photo('B05_DC11_A6_Amplitude-3p480V.png',12.8,'Hình 6. DC11.A6: xung mức thấp gần 0 V, mức cao gần +3,5 V; Amplitude trên ảnh 3,480 V.')
photo('B05_DC11_A7-vang_A8-xanh_Dong-thoi_400Hz.png',12.8,'Hình 7. Đo đồng thời A7 màu vàng và A8 màu xanh; cả hai 400,0 Hz và đảo trạng thái với nhau.')
p('A7 và A8 có Amplitude lần lượt 26,27 V và 26,26 V, phù hợp khoảng chênh hai rail điều khiển. Khi A7 cao, A8 thấp và ngược lại ở phần ổn định. Chu kỳ tương ứng là 2,50 ms. Đây là bằng chứng trực tiếp cho chức năng tạo hai tín hiệu điều khiển bù nhau.')
p('B05 được người thực hiện chấp nhận đạt. Ảnh toàn chu kỳ không định lượng chồng dẫn rất ngắn tại sườn; số Frequency tự động 351,3 kHz ở ảnh A6 không đại diện chu kỳ xung chính. Không sử dụng các số đọc cursor sai khoảng thời gian để suy tần số hoặc độ lệch pha.')

# 7
page(); h('6 Kết quả thử đầu ra có tải')
p('B10 đã được thực hiện lần lượt trên các kênh. Hai ảnh dưới là kết quả đại diện tại mức vào +1 V và +4 V. Khi cấp −1 V và −4 V, người thực hiện ghi nhận dạng sóng tương tự nhưng đảo lại, phù hợp cơ chế biểu diễn dấu bằng pha của tấm.')
photo('B10_OUT_Vin-plus1V_Amplitude-1p975V_Kenh-chua-ghi.png',12.6,'Hình 8. Vin = +1 V DC: tín hiệu ra hai cực tính, Amplitude CH2 = 1,975 V; cursor khoảng −1,082/+1,011 V.')
photo('B10_OUT_Vin-plus4V_Amplitude-7p957V_Kenh-chua-ghi.png',12.6,'Hình 9. Vin = +4 V DC: Amplitude CH2 = 7,957 V; cursor khoảng +3,970/−3,821 V; Frequency hiển thị 400,0 Hz.')
p('Các ảnh cho thấy đầu ra có hai mức gần ±Vin, tăng biên độ theo mức vào và không thấy mất một nửa chu kỳ hay ghim hẳn vào rail ở thang quan sát. Ảnh đại diện chưa gắn mã kênh cụ thể; các kết quả ở mức âm được ghi theo xác nhận của người thực hiện, không có ảnh riêng trong hồ sơ này.')

# 8
page(); h('7 Tổng hợp kết quả và đánh giá kỹ thuật')
p('B07 kiểm tra tĩnh kênh 1a đã xác nhận hai đường hoạt động: pha cùng dấu cho S và Vo gần +1 V; pha đảo cho S gần 0 V và Vo gần −1 V khi Vin vẫn +1 V. Kết quả phù hợp với quan hệ Vo ≈ 2S − Vin, giúp liên hệ được trạng thái khóa với điện áp đầu ra thay vì chỉ đánh giá bằng cảm quan dạng sóng.')
table(['Nhóm kênh','B09 không tải','B10 có tải'],[
('1a và 1b','Đã thử; mức ra tương đồng','Đã thử; không báo bất thường'),
('2a và 2b','Đã thử; mức ra tương đồng','Đã thử; không báo bất thường'),
('3a và 3b','Đã thử; mức ra tương đồng','Đã thử; không báo bất thường'),
('4a và 4b','Đã thử; mức ra tương đồng','Đã thử; không báo bất thường'),
('5a và 5b','Đã thử; mức ra tương đồng','Đã thử; không báo bất thường'),
('6a và 6b','Đã thử; mức ra tương đồng','Đã thử; không báo bất thường'),
('7a và 7b','Đã thử; mức ra tương đồng','Đã thử; không báo bất thường'),
('8a và 8b','Đã thử; mức ra tương đồng','Đã thử; không báo bất thường')
],[3.1,7,7])
p('Bảng tổng hợp phạm vi theo xác nhận đã kiểm tra toàn bộ kênh của người thực hiện; không phải bảng số đo riêng từng kênh. B09 ghi nhận các mức điện áp ra gần giống nhau. B10 ghi nhận đáp ứng ở mức dương và âm; hai ảnh tại trang trước minh họa kết quả có tải.')
sub('Ý nghĩa của kết quả có tải')
p('Mức Amplitude 1,975 V ở đầu vào +1 V và 7,957 V ở +4 V gần mức chênh danh định 2 V và 8 V. Sự thay đổi này hỗ trợ đánh giá đáp ứng theo biên độ đầu vào. Chưa dùng các con số trên để suy sai số gain hoặc tuyến tính chính xác vì chưa có Vin thực đo đồng thời, bảng số riêng từng kênh và thiết lập đo biên độ thống nhất.')
sub('Các kiểm tra mở rộng')
p('B11 kiểm tra ảnh hưởng sang kênh khác và B12 kiểm tra điều kiện mở rộng đã thực hiện ở một số điểm, chưa phát hiện vấn đề. Kết quả này bổ sung cho đánh giá chức năng, nhưng chưa chứng minh toàn bộ tổ hợp tải, hoạt động đồng thời tất cả kênh hoặc độ ổn định dài hạn. Loại tải thực và thời lượng B12 chưa được ghi thành số trong biên bản.')

# 9
page(); h('8 Kết luận và đề xuất xử lý')
p('Kết luận: tấm A313-V.06 hoạt động tốt về chức năng trong phạm vi kiểm tra đã thực hiện.',True)
p('Kết luận dựa trên chuỗi bằng chứng nhất quán: có tín hiệu tham chiếu sau chia áp; U14/U15 tạo chuyển mức; M và M̄ cùng 400 Hz và bù trạng thái; phép giữ pha kênh 1a xác nhận đường truyền cùng dấu và đảo dấu; thử động và tải trên các kênh cho mức ra tương đồng, thay đổi phù hợp với mức và dấu đầu vào. Các điểm kiểm tra chọn mẫu B11–B12 chưa ghi nhận bất thường.')
p('Chưa có căn cứ từ các phép thử hiện tại để chỉ định linh kiện hỏng hoặc thay thế linh kiện trên tấm. Đề nghị ghi nhận kết quả kiểm tra chức năng đạt và giữ nguyên cấu hình linh kiện đã thử. Khi đưa trở lại khối, thực hiện kiểm tra tích hợp theo hướng dẫn khai thác tương ứng để xác nhận đồng thời giao tiếp với các tấm liên quan và tải ngoại vi. [1]')
sub('Giá trị của hồ sơ kỹ thuật đã hình thành')
p('Hồ sơ đã liên kết được ba lớp thông tin: vai trò của tấm trong hệ thống, đường truyền tới từng linh kiện trên sơ đồ và phép đo thực tế tại các nút chức năng. Bảng ánh xạ 16 kênh, cách giữ hai pha DC và bộ ảnh có tên điểm đo tạo cơ sở tái kiểm tra hoặc khoanh vùng nhanh nếu xuất hiện lỗi về sau.')
p('Các kết quả được tách rõ thành số đọc trực tiếp trên ảnh, kết quả xác nhận của người thực hiện và giá trị danh định phục vụ phân tích. Cách ghi này giúp bảo toàn khả năng đối chiếu, tránh biến số dự kiến thành số thực đo hoặc dùng một ảnh đại diện thay cho toàn bộ bảng sai số kênh.')
sub('Nội dung cần hoàn thiện khi nghiệm thu đầy đủ')
p('Bổ sung bảng Vin/Vo thực từng kênh để lượng hóa gain và offset nếu yêu cầu nghiệm thu quy định; ghi mã kênh, trị số tải và điều kiện đo cho từng ảnh; hoàn thiện ma trận tải thực và kiểm tra tích hợp. Đây là phần hoàn thiện hồ sơ định lượng, không phải dấu hiệu đã phát hiện hỏng trong đợt thử chức năng này.')
sub('Tài liệu và bằng chứng đối chiếu')
p('[1] HDSD-A313-V_HS dau A.pdf, mục 1.2 và 3.2.4, trang in 4–5 và 10–11; mục sửa chữa và kiểm tra lại thiết bị.\n[2] BVTK-A313-V.06.pdf: BV.22 sơ đồ nguyên lý, BV.23–BV.24 bố trí mạch và BV.25 linh kiện.\n[3] PHAN-TICH_A313-V.06_KHUECH-DAI-TIN-HIEU-RA.md, Rev.C: đối chiếu net, ánh xạ 16 kênh, cơ chế điều chế và phương pháp thử.\n[4] QUY-TRINH_TEST_TUNG-BUOC_A313-V.06.md và nhật ký thực hiện; ảnh gốc trong KET-QUA-TEST_A313-V.06.')
p('Người lập báo cáo                                   Người kiểm tra\n\n\n....................................                         ....................................')

assert len(d.inline_shapes)==11, len(d.inline_shapes)
assert len(d.tables)==5, len(d.tables)
for st in d.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    fonts=d.styles[name].element.get_or_add_rPr().rFonts
    for key in list(fonts.attrib):
        if 'theme' in key.lower(): del fonts.attrib[key]
    for key in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+key),'Times New Roman')
d.save(OUT)
print(OUT)
